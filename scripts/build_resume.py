"""Build the downloadable resume from Kerk's original resume template."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import hashlib
import shutil
import sys
import tempfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DEFAULT_TEMPLATE = Path("/Users/KERK_Zhi_Sheng/Downloads/Resume - Kerk Zhi Sheng.docx")
EXPECTED_TEMPLATE_SHA256 = "dac8d03db16337fc2a712d4b21697bedb7641ab6476f0d0202f11cbeea9e702f"
DEFAULT_OUTPUT = Path("public/Kerk_Zhi_Sheng_Resume.docx")
PORTFOLIO_URL = "https://kerk-builds.vercel.app"
PRESERVE_PARTS = {
    "[Content_Types].xml",
    "_rels/.rels",
    "customXML/_rels/item1.xml.rels",
    "customXML/item1.xml",
    "customXML/itemProps1.xml",
    "word/_rels/fontTable.xml.rels",
    "word/fontTable.xml",
    "word/fonts/NotoSansSymbols-bold.ttf",
    "word/fonts/NotoSansSymbols-regular.ttf",
    "word/footer1.xml",
    "word/footer2.xml",
    "word/footer3.xml",
    "word/header1.xml",
    "word/header2.xml",
    "word/header3.xml",
    "word/media/image1.png",
    "word/numbering.xml",
    "word/settings.xml",
    "word/styles.xml",
    "word/theme/theme1.xml",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clear_after_header(doc: Document) -> None:
    """Keep the source name, rule, contact row, and spacer; replace the body."""
    body = doc._element.body
    section_properties = body.sectPr
    children = list(body)
    for child in children[3:]:
        if child is not section_properties:
            body.remove(child)


def set_run(run, *, bold: bool | None = None, size: float = 11.2) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def add_hyperlink(paragraph, text: str, url: str, *, size: float = 10.2) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), str(round(size * 2)))
    properties.extend([fonts, color, underline, font_size])
    run.append(properties)

    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_body_paragraph(
    paragraph,
    *,
    before: float = 0,
    after: float = 3,
    keep_next: bool = False,
) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.widow_control = True


def add_section_heading(doc: Document, label: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(label.upper())
    set_run(run, bold=True, size=11.5)
    run.underline = True


def set_cell_margins_zero(cell) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")


def set_table_borders_none(table) -> None:
    table_properties = table._tbl.tblPr
    borders = table_properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def add_header_row(doc: Document, label: str, dates: str, *, before: float = 7) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(5.45)
    table.columns[1].width = Inches(1.82)
    set_table_borders_none(table)

    left_cell, right_cell = table.rows[0].cells
    for cell in (left_cell, right_cell):
        set_cell_margins_zero(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM

    left = left_cell.paragraphs[0]
    set_body_paragraph(left, before=before, after=3, keep_next=True)
    set_run(left.add_run(label), bold=True, size=11.2)

    right = right_cell.paragraphs[0]
    set_body_paragraph(right, before=before, after=3, keep_next=True)
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(right.add_run(dates), bold=True, size=11.2)


def add_item_heading(doc: Document, label: str, *, before: float = 7) -> None:
    paragraph = doc.add_paragraph()
    set_body_paragraph(paragraph, before=before, after=3, keep_next=True)
    set_run(paragraph.add_run(label), bold=True, size=11.2)


def add_plain_paragraph(doc: Document, text: str, *, after: float = 4) -> None:
    paragraph = doc.add_paragraph()
    set_body_paragraph(paragraph, after=after)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run(paragraph.add_run(text), size=11.0)


def add_subsection_heading(doc: Document, label: str, *, before: float = 5) -> None:
    paragraph = doc.add_paragraph()
    set_body_paragraph(paragraph, before=before, after=3, keep_next=True)
    run = paragraph.add_run(label.upper())
    set_run(run, bold=True, size=9.6)


def add_numbering(paragraph, *, num_id: int = 4, level: int = 0) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    numbering_properties = paragraph_properties.find(qn("w:numPr"))
    if numbering_properties is None:
        numbering_properties = OxmlElement("w:numPr")
        paragraph_properties.insert(0, numbering_properties)
    level_node = OxmlElement("w:ilvl")
    level_node.set(qn("w:val"), str(level))
    num_node = OxmlElement("w:numId")
    num_node.set(qn("w:val"), str(num_id))
    numbering_properties.append(level_node)
    numbering_properties.append(num_node)


def add_bullet(doc: Document, text: str, *, after: float = 3.5, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    set_body_paragraph(paragraph, after=after)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_numbering(paragraph, num_id=4, level=0)
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        set_run(lead, bold=True, size=11.0)
    set_run(paragraph.add_run(text), size=11.0)


def restore_preserve_only_parts(reference: Path, generated: Path) -> None:
    """Restore template-owned package parts byte-for-byte after python-docx saves."""
    with ZipFile(reference) as reference_zip, ZipFile(generated) as generated_zip:
        generated_parts = {name: generated_zip.read(name) for name in generated_zip.namelist()}
        for part in PRESERVE_PARTS:
            generated_parts[part] = reference_zip.read(part)

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as handle:
        temporary_path = Path(handle.name)
    try:
        with ZipFile(temporary_path, "w", ZIP_DEFLATED) as output_zip:
            for name, data in generated_parts.items():
                output_zip.writestr(name, data)
        shutil.move(temporary_path, generated)
    finally:
        temporary_path.unlink(missing_ok=True)


def build_resume(output_path: Path, template_path: Path = DEFAULT_TEMPLATE) -> None:
    if not template_path.exists():
        raise FileNotFoundError(f"Resume template not found: {template_path}")
    template_hash = sha256(template_path)
    if template_hash != EXPECTED_TEMPLATE_SHA256:
        raise ValueError(f"Template changed unexpectedly: {template_hash}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, output_path)
    doc = Document(output_path)
    clear_after_header(doc)

    contact = doc.paragraphs[1]
    for child in list(contact._p):
        if child.tag != qn("w:pPr"):
            contact._p.remove(child)
    contact.paragraph_format.space_after = Pt(3)
    contact.paragraph_format.tab_stops.add_tab_stop(Inches(2.7), WD_TAB_ALIGNMENT.LEFT)
    contact.paragraph_format.tab_stops.add_tab_stop(Inches(4.55), WD_TAB_ALIGNMENT.LEFT)
    set_run(contact.add_run("Email: zhishengkerk@gmail.com\t"), size=10.2)
    set_run(contact.add_run("Mobile: +65 8188 7419\t"), size=10.2)
    set_run(contact.add_run("Portfolio: "), size=10.2)
    add_hyperlink(contact, "kerk-builds.vercel.app", PORTFOLIO_URL)

    add_section_heading(doc, "Profile")
    add_plain_paragraph(
        doc,
        "Data and AI Engineer with experience delivering internal data products across government and semiconductor operations. My work spans production data pipelines, operational monitoring, supply chain optimization, and two AI applications used in production. I build from the data model outward so workflows remain traceable, maintainable, and ready for internal adoption.",
        after=5,
    )

    add_section_heading(doc, "Education")
    add_header_row(doc, "National University of Singapore (NUS)", "Aug 2018 to May 2022")
    add_bullet(doc, "Bachelor of Engineering, Industrial and Systems Engineering")

    add_section_heading(doc, "Professional Experience")
    add_header_row(doc, "GovTech  |  Data and AI Engineer", "2025 to Present", before=2)
    add_plain_paragraph(
        doc,
        "Work across five internal data projects while extending this foundation into AI product engineering. Delivered two internal AI applications into production, with responsibility spanning data design, application workflows, deployment, and operational readiness.",
        after=5,
    )
    add_subsection_heading(doc, "Core data engineering", before=3)
    for bullet in [
        "Support five internal projects by developing operational dashboards and building, maintaining, and optimizing the data pipelines used by internal teams.",
        "Built monitoring solutions for pipeline skip detection and volume analysis, helping teams surface missed runs and unexpected processing patterns for investigation.",
    ]:
        add_bullet(doc, bullet, after=4)

    add_subsection_heading(doc, "AI product engineering", before=5)
    for bullet in [
        "Delivered InboxPilot and BoardFlow, two internal AI applications now used in production, from initial data model and workflow design through deployment.",
        "Built InboxPilot as a channel-agnostic multi-agent triage system that classifies requests, retrieves governed knowledge, scores confidence, and routes each case to an automated reply or human review.",
        "Designed BoardFlow as an AI-enabled workspace that turns conversations into structured ticket proposals and connects them with Kanban workflows, configurable agents, permissions, and cross-workspace operations.",
        "Established intentional schemas, dependable persistence, versioned knowledge, and auditable agent decisions so both products can be operated and extended beyond an initial demonstration.",
        "Developed an internal real-time pose-detection engine with MediaPipe for an internal workflow.",
    ]:
        add_bullet(doc, bullet, after=4)

    add_subsection_heading(doc, "Internal knowledge sharing", before=5)
    for bullet in [
        "Presented an internal Databricks brown-bag session on Genie agents and Genie Code, sharing how internal teams use them to shorten the path from governed data to useful insights.",
        "Participated in internal speaking forums and hackathons to exchange practical AI delivery patterns and test new product ideas with colleagues.",
    ]:
        add_bullet(doc, bullet, after=5)

    add_header_row(doc, "Micron Technology  |  Data Science Engineer", "Jun 2022 to 2025", before=11)
    add_subsection_heading(doc, "Optimization and planning", before=3)
    for bullet in [
        "Led supply-chain optimization work for tactical planning and planned-order firming, including improvements that delivered an additional 2-5% cost saving for assembly products while preserving order constraints.",
        "Built raw-material health reports and automated tactical-planning reports to support timely order fulfillment and downstream firming operations.",
        "Led equipment performance-to-model tracking by comparing actual and modelled equipment time, then coordinated data engineers to establish trusted sources for accuracy analysis.",
    ]:
        add_bullet(doc, bullet, after=4)

    add_subsection_heading(doc, "Data platforms and automation", before=5)
    for bullet in [
        "Built and maintained more than 40 data pipelines and automation flows, together with more than 60 trusted tables for optimization, planning reports, and raw-material health analysis.",
        "Led the Planned Order Firming migration to Snowflake and automated the end-to-end workflow, saving more than 80 hours each week.",
        "Scaled product-assembly re-entrance reporting across product groups, contributing a further 20 hours of weekly time savings.",
    ]:
        add_bullet(doc, bullet, after=4)

    add_header_row(doc, "National University of Singapore  |  Research Intern", "May to Aug 2021", before=11)
    add_subsection_heading(doc, "Research and simulation", before=3)
    for bullet in [
        "Conducted a technology scan of PSA's operating context and assessed technologies with potential strategic value.",
        "Supported a Huawei warehouse-simulation project by refining entity-flow diagrams and defining grid-based path-mover behaviour.",
        "Implemented the XML backbone that translated warehouse process flows into the simulation model.",
    ]:
        add_bullet(doc, bullet, after=4)

    add_section_heading(doc, "Certifications")
    add_bullet(doc, "Databricks Certified Data Engineer Associate")
    add_bullet(doc, "DART AI for Cybersecurity Practitioners")
    add_bullet(doc, "AI Singapore AI for Industry® Literacy in AI")

    add_section_heading(doc, "Skills")
    add_bullet(doc, "LLM orchestration, multi-agent systems, retrieval-augmented generation, and AI agents", bold_lead="AI systems: ")
    add_bullet(doc, "Databricks, Snowflake, PostgreSQL, NiFi, SQL, and ETL", bold_lead="Data platforms: ")
    add_bullet(doc, "TypeScript, React, Next.js, Node.js, and Docker", bold_lead="Product engineering: ")
    add_bullet(doc, "Python, Tableau, R, optimization, and simulation", bold_lead="Analytics: ")

    doc.save(output_path)
    restore_preserve_only_parts(template_path, output_path)


if __name__ == "__main__":
    destination = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUTPUT
    source = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_TEMPLATE
    build_resume(destination, source)
